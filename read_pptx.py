from pptx import Presentation
import sys

def read_pptx(filepath, output_path):
    try:
        prs = Presentation(filepath)
    except Exception as e:
        print(f"Error opening presentation: {e}")
        sys.exit(1)
        
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(f"Presentation has {len(prs.slides)} slides.\n")
        
        for i, slide in enumerate(prs.slides):
            f.write(f"\n--- Slide {i+1} ---\n")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    f.write(f"{shape.text}\n")
                elif shape.shape_type == 6: # GROUP
                    for child in shape.shapes:
                        if child.has_text_frame:
                            f.write(f"{child.text}\n")

if __name__ == "__main__":
    read_pptx(r"C:\Users\Dell\Downloads\Smart_Krishi_Yatra_SKH_Final_Template.pptx", "pptx_content.txt")
