from pptx import Presentation
from pptx.util import Inches, Pt
import sys
import os

def replace_text(shape, replacements):
    if shape.has_text_frame:
        text = shape.text
        for old, new in replacements:
            if old in text:
                shape.text = text.replace(old, new)
                # Quick fix for font size reset by shape.text assignment
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(14) 
    
    if shape.shape_type == 6: # GROUP
        for child in shape.shapes:
            replace_text(child, replacements)

def modify_presentation(template_path, output_path):
    prs = Presentation(template_path)
    
    # Define replacements
    replacements_slide1 = [
        ("Problem Statement ID \xef\xbf\xbd Problem Statement Title- Theme- PS Category- Software/Hardware Team ID- Team",
         "Problem Statement ID: SKH041\nTitle: Smart Krishi-Yatra AI\nTheme: Agriculture\nCategory: Software\nTeam ID: [Your Team ID]\nTeam Name: Team Agnivega")
    ]
    
    replacements_slide2 = [
        ("Proposed Solution (Describe your Idea/Solution/Prototype)   Detailed explanation of the proposed sol",
         "Smart Krishi-Yatra AI is a market-aware agricultural logistics OS.\n\nIt determines WHERE, WHEN, and HOW a farmer should transport produce to maximize EXPECTED NET REALIZATION (ENR).\n\nFeatures:\n• Voice-first multilingual input (Marathi, Hindi).\n• Deterministic CVRP-based optimizer for load pooling.\n• Dynamic Spoilage & Queue tracking.")
    ]
    
    replacements_slide3 = [
        ("Technologies to be used (e.g. programming languages, frameworks, hardware) Methodology and process f",
         "Technologies & Architecture:\n• Frontend: React 19, TypeScript, TanStack PWA\n• Styling: Tailwind CSS, Radix UI\n• Routing Engine: Node.js CVRP Optimizer\n• Geographic Data: OSRM API (Primary) / Haversine (Offline Fallback)\n\nData Structures:\n• Graph-based distance matrices for CVRP\n• Priority queues for load optimization")
    ]
    
    replacements_slide4 = [
        ("Analysis of the feasibility of the idea Potential challenges and risks Strategies for overcoming the",
         "Feasibility: Fully feasible as a web/PWA platform without requiring native app store downloads.\n\nChallenges & Risks:\n• Low digital literacy in rural areas.\n• Unstable internet connections.\n\nStrategies:\n• Implemented Voice-first IVR-style interactions.\n• Tiered offline-capable architecture (Local calculations via CVRP heuristics).")
    ]
    
    replacements_slide5 = [
        ("Potential impact on the target audience Benefits of the solution (social, economic, environmental, e",
         "Impact & Benefits:\n• Economic: Farmers maximize Expected Net Realization (ENR), preventing unprofitable trips.\n• Social: Encourages community pooling, distributing transport costs.\n• Environmental: Fewer half-empty trucks on the road due to CVRP load aggregation.\n• Agricultural: Prevents post-harvest loss through real-time transit risk modeling.")
    ]
    
    replacements_slide6 = [
        ("Details / Links of the reference and research work",
         "References & Research:\n1. Capacitated Vehicle Routing Problem (CVRP) heuristics.\n2. OSRM (Open Source Routing Machine).\n3. Govt. APMC / Mandi pricing models.\n4. GitHub Repository: github.com/takshalchaudhari/AgniVega")
    ]

    replacements_general = [
        ("Problem Statement ID \xef\xbf\xbd", "Problem Statement ID: SKH041"),
        ("Problem Statement Title-", "Title: Smart Krishi-Yatra AI"),
        ("Theme-", "Theme: Agriculture"),
        ("PS Category- Software/Hardware", "Category: Software"),
        ("Team ID-", "Team ID: [Your Team ID]"),
        ("Team Name \xef\xbf\xbd", "Team Name: Team Agnivega")
    ]

    for i, slide in enumerate(prs.slides):
        reps = replacements_general.copy()
        if i == 0: reps += replacements_slide1
        elif i == 1: reps += replacements_slide2
        elif i == 2: reps += replacements_slide3
        elif i == 3: reps += replacements_slide4
        elif i == 4: reps += replacements_slide5
        elif i == 5: reps += replacements_slide6

        # Try replacing exactly the whole text block if there are weird invisible characters
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text
                if i == 0 and "Problem Statement ID" in text:
                    shape.text = "Problem Statement ID: SKH041\nTitle: Smart Krishi-Yatra AI\nTheme: Agriculture\nCategory: Software\nTeam ID: [Your Team ID]\nTeam Name: Team Agnivega"
                elif i == 1 and "Proposed Solution" in text:
                    shape.text = replacements_slide2[0][1]
                elif i == 2 and "Technologies to be" in text:
                    shape.text = replacements_slide3[0][1]
                elif i == 3 and "Analysis of the feasibility" in text:
                    shape.text = replacements_slide4[0][1]
                elif i == 4 and "Potential impact" in text:
                    shape.text = replacements_slide5[0][1]
                elif i == 5 and "reference and research work" in text:
                    shape.text = replacements_slide6[0][1]
                
                # Fix font size
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(16)
            
            replace_text(shape, reps)

    # Add images
    banner_path = r"C:\Users\Dell\Downloads\smart-krishi-yatra-ai (1)\public\assets\readme_banner_3d.jpg"
    arch_path = r"C:\Users\Dell\.gemini\antigravity\brain\5a6866de-93c8-441e-a5f4-55252663b455\skh_architecture_diagram_1786389054927.jpg"
    
    if os.path.exists(banner_path):
        # Add to slide 2
        prs.slides[1].shapes.add_picture(banner_path, Inches(5), Inches(2), width=Inches(4))
        
    if os.path.exists(arch_path):
        # Add to slide 3
        prs.slides[2].shapes.add_picture(arch_path, Inches(5), Inches(2), width=Inches(4))
        
    # Delete slide 7 if exists (Important Instructions)
    if len(prs.slides) > 6:
        xml_slides = prs.slides._sldIdLst
        slides = list(xml_slides)
        xml_slides.remove(slides[6])

    prs.save(output_path)
    print(f"Saved successfully to {output_path}")

if __name__ == "__main__":
    modify_presentation(r"C:\Users\Dell\Downloads\SKH_IDEA_PRESENTATION_FORMAT.pptx", r"C:\Users\Dell\Downloads\SKH_IDEA_PRESENTATION_Agnivega.pptx")
