import collections 
import collections.abc
from pptx import Presentation
import sys

def inspect_pptx(filepath):
    try:
        prs = Presentation(filepath)
    except Exception as e:
        print(f"Error opening presentation: {e}")
        sys.exit(1)
        
    print(f"Presentation has {len(prs.slides)} slides.")
    
    for i, slide in enumerate(prs.slides):
        print(f"\n--- Slide {i+1} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                print(f"Shape: {shape.name} | Text: {shape.text.replace(chr(10), ' ').replace(chr(13), '')[:100]}")
            else:
                print(f"Shape: {shape.name} | Type: {shape.shape_type}")
            if shape.is_placeholder:
                print(f"  --> Placeholder ID: {shape.placeholder_format.idx}, Type: {shape.placeholder_format.type}")

if __name__ == "__main__":
    inspect_pptx(r"C:\Users\Dell\Downloads\SKH_IDEA_PRESENTATION_FORMAT.pptx")
